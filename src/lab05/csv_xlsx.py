import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

ENCODING = "utf-8"
TITLE = "Sheet1"
MIN_COL_WIDTH = 8

def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """
    Convert CSV into XLSX.

    Uses openpysxl.
    First string of CSV — header.
    Sheet named "Sheet1".
    Columns — auto-width based on text length (at least 8 chars).
    """

    wb = Workbook()
    ws = wb.active

    if ws is None:
        raise RuntimeError("Workbook has no active worksheet")

    ws.title = TITLE

    column_widths = dict()

    csv_path_path = Path(csv_path)
    with csv_path_path.open(encoding=ENCODING) as f:
        for row in csv.reader(f):
            ws.append(row)

            # Find width for each column
            for i, s in enumerate(row):
                column_widths[i+1] = max( # column indexes start with 1
                    column_widths.get(i+1, MIN_COL_WIDTH),
                    len(s))

    for i in column_widths:
        print(column_widths[i])
        ws.column_dimensions[get_column_letter(i)].width = column_widths[i]

    xlsx_path_path = Path(xlsx_path)
    wb.save(xlsx_path_path)
